"""
Tests unitaires du parseur Excel.

La fixture reproduit la structure réelle des classeurs de
`07_Projets_DSIO/Projet encours`, telle que constatée sur les 21 fichiers :
bloc de titre au-dessus du tableau, feuille de tableau de bord en second,
colonne « Projets » fusionnée portant la phase, lignes de données espacées
d'une ligne sur deux, pourcentages au format Excel, lignes de section sans
responsable, et numéros saisis de façon irrégulière.

Ces tests ne nécessitent ni base de données ni stack démarrée.
"""
from datetime import date

import openpyxl
import pytest

from app.services.excel_parser import (
    _parse_responsables,
    _safe_date,
    canonical_numero,
    normalize_numero,
    parse_excel_file,
    parse_workbook,
    split_phase,
)

# Ligne d'en-tête volontairement basse : les fichiers réels la placent entre
# la ligne 11 et la ligne 19.
ENTETE = 15
ENTETES = [
    "Projets", "N° d'actions", "Actions", "Resp. réalisation", "Resp. Suivi",
    "%Progress", "SPI", "OTD", "Deadline", "Date de Réalisation",
    "Charges (h/j)", "Commentaires",
]


def _classeur(chemin, lignes, *, avec_tdb=True, entete=ENTETE, phases=None):
    """Construit un classeur au format DSIO.

    `lignes` : liste de dicts décrivant les lignes de données, dans l'ordre.
    Chaque ligne est écrite toutes les deux lignes, comme dans les fichiers
    réels où les cellules sont fusionnées verticalement.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "P01 - Cantine"

    # Bloc de titre : nom du projet, équipe, légende KPI.
    ws["A1"] = "P01 - Cantine"
    ws["E1"] = "KPI  PROJET"
    ws["A4"] = "Equipe Projet"
    ws["A5"] = "Meylis "
    ws["C5"] = "Technicien informatique"
    ws["A14"] = "Planning Projet SI"

    for index, libelle in enumerate(ENTETES, start=1):
        ws.cell(row=entete, column=index, value=libelle)

    ligne = entete + 1
    for donnees in lignes:
        ws.cell(row=ligne, column=1, value=donnees.get("projets"))
        ws.cell(row=ligne, column=2, value=donnees.get("numero"))
        ws.cell(row=ligne, column=3, value=donnees.get("description"))
        ws.cell(row=ligne, column=4, value=donnees.get("responsables"))
        ws.cell(row=ligne, column=5, value=donnees.get("resp_suivi"))
        for colonne, cle in ((6, "progress"), (7, "spi"), (8, "otd")):
            if cle in donnees:
                cellule = ws.cell(row=ligne, column=colonne, value=donnees[cle])
                if donnees.get("format_pourcent", True):
                    cellule.number_format = "0%"
        ws.cell(row=ligne, column=9, value=donnees.get("deadline"))
        ws.cell(row=ligne, column=10, value=donnees.get("date_realisation"))
        ws.cell(row=ligne, column=11, value=donnees.get("charges_hj"))
        ws.cell(row=ligne, column=12, value=donnees.get("commentaire"))
        ligne += 2

    # Fusion verticale de la colonne « Projets » sur chaque bloc de phase.
    for debut, fin in (phases or []):
        ws.merge_cells(start_row=debut, start_column=1, end_row=fin, end_column=1)

    if avec_tdb:
        # Toujours présent dans les fichiers réels, et jamais une source
        # d'actions.
        tdb = wb.create_sheet("TDB Projet(P01 - Cantine)")
        tdb["A1"] = "Tableau de bord"

    wb.save(chemin)
    return chemin


# --- Normalisation élémentaire ---------------------------------------------

def test_normalize_numero_tolere_les_espaces():
    assert normalize_numero("P01 -02- 05") == "P01-02-05"
    assert normalize_numero("  P01-01  ") == "P01-01"
    assert normalize_numero(None) == ""


def test_split_phase():
    assert split_phase("P01-02-05") == "02"
    assert split_phase("P10-1-1-01") == "1-1"
    assert split_phase("P04-05") is None


def test_canonical_numero_ignore_le_prefixe_saisi():
    """Dans P01, la phase 2 est numérotée « P02 - 01 » : conserver ce préfixe
    produirait un numéro renvoyant à un autre projet."""
    assert canonical_numero("P01", "02", "P02 - 01") == "P01-02-01"
    assert canonical_numero("P04", None, "P04 - 07") == "P04-07"
    assert canonical_numero("P10", "1-2", "P02 -13 - 5") == "P10-1-2-5"


def test_parse_responsables_separateurs():
    assert _parse_responsables("Meylis / Xavier") == ["Meylis", "Xavier"]
    assert _parse_responsables("Andry II, Xavier") == ["Andry II", "Xavier"]
    assert _parse_responsables("AndryII - Teknet") == ["AndryII", "Teknet"]
    assert _parse_responsables("Xavier et Manda") == ["Xavier", "Manda"]
    assert _parse_responsables("Meylis\n") == ["Meylis"]
    assert _parse_responsables(None) == []


@pytest.mark.parametrize(
    "cellule,attendu",
    [
        ("Meylis-Hassen", ["Meylis", "Hassen"]),
        ("Xavier -Hassen", ["Xavier", "Hassen"]),
        ("Xavier- Hassen", ["Xavier", "Hassen"]),
        ("Manoa -Karine", ["Manoa", "Karine"]),
        ("Xavier-Hassen-MDC", ["Xavier", "Hassen", "MDC"]),
        ("Xavier -manoa -karine", ["Xavier", "manoa", "karine"]),
    ],
)
def test_le_tiret_separe_deux_personnes(cellule, attendu):
    """Sept cellules des classeurs mettaient plusieurs personnes derrière un
    tiret sans espacement régulier. Elles produisaient une fiche composite par
    combinaison : 21 actions invisibles du plan de charge de chacun, et
    impossibles à relancer puisqu'une fiche composite n'a pas d'email."""
    assert _parse_responsables(cellule) == attendu


def test_nom_compose_protege(monkeypatch):
    """Un prénom composé serait découpé à tort : `RESPONSABLES_INSECABLES`
    permet de l'exclure sans toucher au code."""
    monkeypatch.setenv("RESPONSABLES_INSECABLES", "Jean-Pierre, Marie-Claire")
    assert _parse_responsables("Jean-Pierre") == ["Jean-Pierre"]
    assert _parse_responsables("Marie-Claire") == ["Marie-Claire"]
    # Les autres restent découpés.
    assert _parse_responsables("Meylis-Hassen") == ["Meylis", "Hassen"]


def test_ou_n_est_pas_un_separateur():
    """« Teknet ou autre » désigne une incertitude, pas deux personnes."""
    assert _parse_responsables("Teknet ou autre") == ["Teknet ou autre"]
    assert _parse_responsables("Equipe projet") == ["Equipe projet"]


def test_parse_responsables_dedoublonne_sans_tenir_compte_de_la_casse():
    """« xavier » et « Xavier » désignent la même personne : deux fiches
    responsable, ce serait deux relances pour la même action."""
    assert _parse_responsables("Xavier / xavier") == ["Xavier"]


@pytest.mark.parametrize(
    "valeur,attendu",
    [
        ("11/6/2023", date(2023, 11, 6)),
        ("2026-06-18", date(2026, 6, 18)),
        ("-", None),          # tiret de remplissage, fréquent dans les fichiers
        ("", None),
        ("120/07/2026", None),  # faute de frappe relevée dans P26
    ],
)
def test_safe_date(valeur, attendu):
    assert _safe_date(valeur) == attendu


# --- Lecture d'un classeur --------------------------------------------------

def _projet_simple(tmp_path):
    return _classeur(
        tmp_path / "p04.xlsx",
        [
            {
                "projets": "P04 - Journalisation", "numero": "P04 - 01",
                "description": "Définir les besoins fonctionnels",
                "responsables": "Meylis / Xavier", "resp_suivi": "Hassen",
                "progress": 1.0, "spi": 1.0, "otd": 1.0,
                "deadline": date(2026, 2, 10), "date_realisation": date(2026, 2, 12),
                "charges_hj": 1.5, "commentaire": "RAS",
            },
            {
                "numero": "P04-02", "description": "Rédiger le cahier des charges",
                "responsables": "Meylis", "progress": 0.5,
                "deadline": date(2026, 12, 31),
            },
        ],
    )


def test_entete_detectee_quel_que_soit_son_rang(tmp_path):
    resultat = parse_workbook(_projet_simple(tmp_path), project_code="P04")
    assert resultat.header_row == ENTETE
    assert resultat.sheet_name == "P01 - Cantine"
    assert len(resultat.actions) == 2


def test_pourcentages_normalises(tmp_path):
    """Une cellule affichée « 100 % » vaut 1.0 dans le fichier : lue telle
    quelle, une action terminée arrivait à 1 % et restait relancée."""
    actions = parse_excel_file(_projet_simple(tmp_path), project_code="P04")
    assert actions[0]["progress"] == 100.0
    assert actions[0]["spi"] == 100.0
    assert actions[1]["progress"] == 50.0


def test_valeurs_deja_en_centiemes_inchangees(tmp_path):
    chemin = _classeur(
        tmp_path / "brut.xlsx",
        [{
            "numero": "P04-01", "description": "Action", "responsables": "Meylis",
            "progress": 100.0, "deadline": date(2026, 1, 1), "format_pourcent": False,
        }],
    )
    assert parse_excel_file(chemin, project_code="P04")[0]["progress"] == 100.0


def test_champs_lus_correctement(tmp_path):
    action = parse_workbook(_projet_simple(tmp_path), project_code="P04").actions[0]
    assert action.numero == "P04-01"
    assert action.description == "Définir les besoins fonctionnels"
    assert action.responsable_names == ["Meylis", "Xavier"]
    assert action.resp_suivi == "Hassen"
    assert action.deadline == date(2026, 2, 10)
    assert action.date_realisation == date(2026, 2, 12)
    assert action.charges_hj == 1.5
    assert action.commentaire == "RAS"


def test_feuille_tableau_de_bord_ignoree(tmp_path):
    """Chaque classeur embarque une feuille `TDB Projet(...)` ; certains
    contiennent aussi d'anciennes copies du planning."""
    resultat = parse_workbook(_projet_simple(tmp_path), project_code="P04")
    assert not resultat.sheet_name.startswith("TDB")


def test_nom_de_feuille_non_fiable(tmp_path):
    """Le classeur de P32 a conservé le nom de feuille de P29 dont il est la
    copie : la sélection doit se faire sur le contenu, pas sur le nom."""
    chemin = _projet_simple(tmp_path)
    wb = openpyxl.load_workbook(chemin)
    wb["P01 - Cantine"].title = "P29 - Réaménagement du réseau"
    wb.save(chemin)
    assert len(parse_workbook(chemin, project_code="P04").actions) == 2


# --- Phases -----------------------------------------------------------------

def test_phase_lue_dans_la_colonne_fusionnee(tmp_path):
    """La phase n'est écrite que sur la première ligne de son bloc : sans
    résolution des fusions, toutes les actions suivantes la perdent."""
    chemin = _classeur(
        tmp_path / "p01.xlsx",
        [
            {"projets": "P01 - Cantine - Phase 1", "numero": "P01 - 01",
             "description": "Etablir le cahier de charges", "responsables": "Meylis",
             "progress": 1.0, "deadline": date(2023, 11, 6)},
            {"numero": "P01 - 02", "description": "Valider le cahier",
             "responsables": "Manda", "progress": 1.0, "deadline": date(2024, 2, 28)},
            {"projets": "P01 - Cantine - Phase 2", "numero": "P02 - 01",
             "description": "Recueillir les besoins", "responsables": "Equipe projet",
             "progress": 1.0, "deadline": date(2026, 2, 10)},
            {"numero": "P02 - 02", "description": "Rédiger le cahier des charges",
             "responsables": "Meylis", "progress": 0.4, "deadline": date(2026, 2, 12)},
        ],
        phases=[(ENTETE + 1, ENTETE + 3), (ENTETE + 5, ENTETE + 7)],
    )
    resultat = parse_workbook(chemin, project_code="P01")
    assert [a.numero for a in resultat.actions] == [
        "P01-01-01", "P01-01-02", "P01-02-01", "P01-02-02"
    ]
    assert resultat.phases == ["01", "02"]


def test_bloc_sans_libelle_rattache_a_la_premiere_phase(tmp_path):
    """Dans P06, seule la phase 2 est annoncée en colonne A ; les premières
    actions doivent malgré tout porter la phase 01, comme dans le fichier
    consolidé."""
    chemin = _classeur(
        tmp_path / "p06.xlsx",
        [
            {"projets": "P06 - Voix sur IP", "numero": "P06 - 01",
             "description": "Analyser l'existant", "responsables": "Xavier",
             "progress": 1.0, "deadline": date(2025, 6, 1)},
            {"projets": "P06 - Voix sur IP - Phase 2", "numero": "P06 - 02",
             "description": "Déployer le trunk SIP", "responsables": "Xavier",
             "progress": 0.2, "deadline": date(2026, 6, 1)},
        ],
    )
    resultat = parse_workbook(chemin, project_code="P06")
    assert [a.numero for a in resultat.actions] == ["P06-01-01", "P06-02-02"]
    assert any("rattachée" in w for w in resultat.warnings)


# --- Lignes de section ------------------------------------------------------

def _projet_hierarchique(tmp_path):
    """Reproduit le découpage de P10 : sections à deux niveaux, numérotation
    incohérente (« P10-12 » pour « P10-1-2 ») et préfixe erroné."""
    return _classeur(
        tmp_path / "p10.xlsx",
        [
            {"numero": "P10-1", "description": "Pointage et contrôle d'accès"},
            {"numero": "P10-1-1", "description": "Pointages Tamatave"},
            {"numero": "P10-1-1 - 01", "description": "Synchroniser les données",
             "responsables": "AndryII - Teknet", "progress": 1.0,
             "deadline": date(2026, 3, 11)},
            {"numero": "P10-12", "description": "Contrôle d'accès"},
            {"numero": "P10-12 - 5", "description": "Valider le système",
             "responsables": "Xavier", "progress": 0.0, "deadline": date(2026, 5, 1)},
            {"numero": "P10-13", "description": "Cantine Tamatave"},
            {"numero": "P02 -13 - 1", "description": "Préparer la connectique",
             "responsables": "Manoa", "progress": 0.0, "deadline": date(2026, 6, 1)},
        ],
    )


def test_lignes_de_section_exclues_des_actions(tmp_path):
    """Une section n'a ni responsable, ni échéance, ni avancement : la
    compter comme une action gonflerait tous les indicateurs."""
    resultat = parse_workbook(_projet_hierarchique(tmp_path), project_code="P10")
    assert resultat.section_rows == 4
    assert len(resultat.actions) == 3
    assert all(a.responsable_names for a in resultat.actions)


def test_niveau_de_section_normalise(tmp_path):
    """« P10-12 » sous « P10-1 » désigne le niveau 1-2 : c'est la forme que
    retient le fichier consolidé."""
    resultat = parse_workbook(_projet_hierarchique(tmp_path), project_code="P10")
    assert [a.numero for a in resultat.actions] == [
        "P10-1-1-01", "P10-1-2-5", "P10-1-3-1"
    ]


def test_prefixe_errone_corrige_par_la_section(tmp_path):
    """L'action « P02 -13 - 1 » se trouve dans le classeur de P10 : le préfixe
    saisi désigne un autre projet et ne doit pas être conservé."""
    resultat = parse_workbook(_projet_hierarchique(tmp_path), project_code="P10")
    derniere = resultat.actions[-1]
    assert derniere.numero == "P10-1-3-1"
    assert derniere.numero_source == "P02-13-1"
    assert derniere.section == "Cantine Tamatave"


# --- Diagnostic -------------------------------------------------------------

def test_tableau_vide_signale(tmp_path):
    """Le classeur de P32 contient un tableau sans aucune ligne remplie :
    renvoyer zéro action sans rien dire laisserait croire à un projet sans
    travaux."""
    chemin = _classeur(tmp_path / "p32.xlsx", [])
    resultat = parse_workbook(chemin, project_code="P32")
    assert resultat.actions == []
    assert any("aucune ligne d'action" in w for w in resultat.warnings)


def test_action_sans_description_signalee(tmp_path):
    chemin = _classeur(
        tmp_path / "p01.xlsx",
        [{"numero": "P01 - 05", "description": None, "responsables": "Meylis",
          "progress": 0.0, "deadline": date(2026, 1, 1)}],
    )
    resultat = parse_workbook(chemin, project_code="P01")
    assert resultat.actions == []
    assert any("sans description" in w for w in resultat.warnings)


def test_echeance_illisible_signalee(tmp_path):
    """Une date illisible sort l'action de toutes les vues de retard : le
    silence serait pire que le bruit."""
    chemin = _classeur(
        tmp_path / "p26.xlsx",
        [{"numero": "P26-11", "description": "Mettre en production",
          "responsables": "Xavier", "progress": 0.0, "deadline": "120/07/2026"}],
    )
    resultat = parse_workbook(chemin, project_code="P26")
    assert resultat.actions[0].deadline is None
    assert any("échéance illisible" in w for w in resultat.warnings)


def test_doublon_de_numero_desambigue_sans_perte(tmp_path):
    """Dans P31, une section numérote dix actions 01, 02, 02, 03, 03… : sans
    désambiguïsation, quatre actions bien réelles étaient écrasées."""
    chemin = _classeur(
        tmp_path / "dbl.xlsx",
        [
            {"numero": "P04-01", "description": "Première", "responsables": "A",
             "progress": 0.0, "deadline": date(2026, 1, 1)},
            {"numero": "P04 - 01", "description": "Doublon", "responsables": "B",
             "progress": 0.0, "deadline": date(2026, 1, 2)},
            {"numero": "P04-01", "description": "Triplon", "responsables": "C",
             "progress": 0.0, "deadline": date(2026, 1, 3)},
        ],
    )
    resultat = parse_workbook(chemin, project_code="P04")
    assert [a.numero for a in resultat.actions] == ["P04-01", "P04-01b", "P04-01c"]
    assert [a.description for a in resultat.actions] == ["Première", "Doublon", "Triplon"]
    assert sum("double" in w for w in resultat.warnings) == 2


def test_code_projet_deduit_si_absent(tmp_path):
    resultat = parse_workbook(_projet_simple(tmp_path))
    assert resultat.project_code == "P04"
    assert any("déduit" in w for w in resultat.warnings)


def test_format_non_supporte(tmp_path):
    chemin = tmp_path / "fichier.csv"
    chemin.write_text("a,b,c", encoding="utf-8")
    with pytest.raises(ValueError):
        parse_workbook(chemin)


def test_fichier_absent(tmp_path):
    with pytest.raises(FileNotFoundError):
        parse_workbook(tmp_path / "inexistant.xlsx")
